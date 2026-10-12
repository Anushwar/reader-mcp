"""End-to-end smoke tests: launch the real server and drive it through an MCP client.

Fixtures are generated at runtime; no sample documents are committed.

Test the published package instead of the checkout with:
    MCP_CMD="uvx --from git+https://github.com/Anushwar/reader-mcp reader" uv run pytest
"""

from __future__ import annotations

import base64
import io
import os
import shlex
import socket
import subprocess
import sys
import time
from pathlib import Path
from typing import Any

import httpx2
import pytest
from docx import Document
from mcp import Client
from mcp.client.stdio import StdioServerParameters
from mcp.client.streamable_http import streamable_http_client
from pypdf import PdfWriter

SERVER_CMD = shlex.split(os.environ.get("MCP_CMD", "")) or [sys.executable, "-m", "reader_mcp"]


@pytest.fixture
def anyio_backend() -> str:
    return "asyncio"


# --------------------------------------------------------------------------- fixtures


def make_pdf(pages: list[str], title: str) -> bytes:
    """Write a minimal text PDF by hand (no PDF-writing dependency needed)."""
    esc = lambda s: s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")  # noqa: E731
    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for pid, text in zip(page_ids, pages):
        stream = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({esc(line)}) Tj T*" for line in text.split("\n")) + " ET"
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       f"/Resources << /Font << /F1 3 0 R >> >> /Contents {pid + 1} 0 R >>".encode())
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())
    objects.append(f"<< /Title ({esc(title)}) /Author (reader tests) >>".encode())

    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info {len(objects)} 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n").encode()
    return out


@pytest.fixture
def workspace(tmp_path: Path) -> dict[str, Path]:
    docs = tmp_path / "docs"
    docs.mkdir()
    (docs / "report.pdf").write_bytes(make_pdf([
        "Quarterly Report\nRevenue grew 12 percent in Q3, driven by enter-\nprise sales.",
        "Risks\nSupply chain delays may affect the Zephyr launch timeline.",
    ], title="Q3 Report"))

    blank = PdfWriter()
    blank.add_blank_page(612, 792)
    buffer = io.BytesIO()
    blank.write(buffer)
    (docs / "scan.pdf").write_bytes(buffer.getvalue())

    doc = Document()
    doc.core_properties.title, doc.core_properties.author = "Team Notes", "Ada"
    doc.add_heading("Overview", level=1)
    doc.add_paragraph("The platform team owns ingestion and search.")
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("Name", "Role"), ("Ada", "Lead"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.save(docs / "notes.docx")

    (docs / "sales.csv").write_text("region,amount,date\nNorth,1200,2026-01-05\nSouth,950,2026-02-11\n")
    (docs / ".private").mkdir()
    (docs / ".private" / "secret.txt").write_text("do not read")
    (tmp_path / "outside.txt").write_text("outside the allowed root")
    return {"root": tmp_path, "docs": docs, "data": tmp_path / "data"}


def server_env(ws: dict[str, Path], **extra: str) -> dict[str, str]:
    env = {k: v for k, v in os.environ.items() if not k.startswith("READER_")}
    return {**env, "READER_DATA_DIR": str(ws["data"]), **extra}


async def call(client: Client, tool: str, **args: Any) -> dict[str, Any]:
    result = await client.call_tool(tool, args)
    if result.is_error:
        return {"error": " ".join(getattr(c, "text", "") for c in result.content)}
    return result.structured_content or {}


# --------------------------------------------------------------------------- stdio


@pytest.mark.anyio
async def test_stdio_end_to_end(workspace: dict[str, Path]) -> None:
    docs = workspace["docs"]
    params = StdioServerParameters(command=SERVER_CMD[0], args=SERVER_CMD[1:],
                                   env=server_env(workspace, READER_ALLOWED_ROOTS=str(docs)))
    async with Client(params) as client:
        tools = {t.name for t in (await client.list_tools()).tools}
        assert tools == {"ingest_document", "list_documents", "get_metadata", "read_document",
                         "search_documents", "delete_document"}

        pdf = await call(client, "ingest_document", path=str(docs / "report.pdf"))
        assert pdf["type"] == "pdf" and pdf["unit_count"] == 2 and not pdf["needs_ocr"]
        assert pdf["properties"]["title"] == "Q3 Report"
        assert "Revenue grew" in pdf["preview"]
        again = await call(client, "ingest_document", path=str(docs / "report.pdf"))
        assert again["doc_id"] == pdf["doc_id"] and again["already_ingested"]

        scan = await call(client, "ingest_document", path=str(docs / "scan.pdf"))
        assert scan["needs_ocr"] is True

        notes = await call(client, "ingest_document", path=str(docs / "notes.docx"))
        assert notes["type"] == "docx" and notes["properties"]["author"] == "Ada"
        assert notes["headings"] == ["Overview"] and notes["table_count"] == 1

        sales = await call(client, "ingest_document", path=str(docs / "sales.csv"))
        assert sales["row_count"] == 2
        assert [(c["name"], c["type"]) for c in sales["columns"]] == [("region", "text"), ("amount", "number"),
                                                                      ("date", "date")]

        # "ﬁ" is the "fi" ligature PDFs often contain; it must stay searchable as plain "fi".
        md = base64.b64encode("# Handbook\nWelcome to the team. We review classiﬁcation errors weekly.".encode()).decode()
        handbook = await call(client, "ingest_document", content_base64=md, filename="handbook.md")
        assert handbook["type"] == "md" and handbook["properties"]["title"] == "Handbook"
        ligature = await call(client, "search_documents", query="classification", doc_id=handbook["doc_id"])
        assert ligature["count"] == 1

        page = await call(client, "read_document", doc_id=pdf["doc_id"], start=1, end=1)
        assert page["items"][0]["page"] == 1 and page["next_start"] == 2

        hits = await call(client, "search_documents", query="Zephyr launch")
        assert hits["results"][0]["doc_id"] == pdf["doc_id"] and hits["results"][0]["page"] == 2
        rejoined = await call(client, "search_documents", query="enterprise", doc_id=pdf["doc_id"])
        assert rejoined["count"] == 1 and rejoined["results"][0]["page"] == 1  # "enter-\nprise" was re-joined
        scoped = await call(client, "search_documents", query="enterprise", doc_id=sales["doc_id"])
        assert scoped["count"] == 0

        meta = await call(client, "get_metadata", doc_id=notes["doc_id"])
        assert meta["filename"] == "notes.docx"
        listed = await call(client, "list_documents")
        assert listed["count"] == 5

        assert (await call(client, "delete_document", doc_id=handbook["doc_id"]))["deleted"]
        assert "Unknown doc_id" in (await call(client, "get_metadata", doc_id=handbook["doc_id"]))["error"]

        # Guard rails
        assert "Hidden" in (await call(client, "ingest_document", path=str(docs / ".private/secret.txt")))["error"]
        assert "outside" in (await call(client, "ingest_document",
                                         path=str(workspace["root"] / "outside.txt")))["error"]
        assert "https" in (await call(client, "ingest_document", url="http://example.com/a.pdf"))["error"]
        assert "private" in (await call(client, "ingest_document", url="https://127.0.0.1/a.pdf"))["error"]
        assert "exactly one" in (await call(client, "ingest_document"))["error"]


# --------------------------------------------------------------------------- http


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@pytest.mark.anyio
async def test_http_upload_flow(workspace: dict[str, Path]) -> None:
    port, token = _free_port(), "test-token"
    proc = subprocess.Popen([*SERVER_CMD, "--http", "--port", str(port)],
                            env=server_env(workspace, READER_AUTH_TOKEN=token, READER_UPLOAD_SECRET="s3cret"),
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        base = f"http://127.0.0.1:{port}"
        for _ in range(100):
            try:
                socket.create_connection(("127.0.0.1", port), timeout=0.2).close()
                break
            except OSError:
                time.sleep(0.1)

        async with httpx2.AsyncClient() as http:
            assert (await http.post(f"{base}/mcp", json={})).status_code == 401

            auth = httpx2.AsyncClient(headers={"Authorization": f"Bearer {token}"})
            async with Client(streamable_http_client(f"{base}/mcp", http_client=auth)) as client:
                tools = {t.name for t in (await client.list_tools()).tools}
                assert "create_upload_url" in tools

                upload = await call(client, "create_upload_url", filename="sales.csv")
                url = upload["upload_url"].replace("http://localhost", "http://127.0.0.1")
                assert (await http.get(url)).status_code == 200  # browser upload page
                assert (await http.put(url + "x", content=b"a,b\n1,2\n")).status_code == 403  # tampered
                put = await http.put(url, content=(workspace["docs"] / "sales.csv").read_bytes())
                assert put.status_code == 200 and put.json()["upload_id"] == upload["upload_id"]
                assert (await http.put(url, content=b"again")).status_code == 409  # single use

                doc = await call(client, "ingest_document", upload_id=upload["upload_id"])
                assert doc["type"] == "csv" and doc["source"] == "upload" and doc["row_count"] == 2
                disabled = await call(client, "ingest_document", path=str(workspace["docs"] / "sales.csv"))
                assert "disabled" in disabled["error"]
            await auth.aclose()
    finally:
        proc.terminate()
        proc.wait(timeout=10)
